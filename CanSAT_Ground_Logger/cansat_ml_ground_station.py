
"""
===============================================================
        CanSAT + ML WEATHER GROUND STATION
===============================================================

Hardware:
    CanSAT
       ↓
    SX1278 LoRa
       ↓
    ESP32-S3 Ground Station
       ↓
    USB Serial
       ↓
    This Python program

Functions:
    1. Receive live CanSAT telemetry
    2. Save telemetry to Excel
    3. Display live graphs
    4. Load LSTM weather model
    5. Load Random Forest weather-condition model
    6. Generate ML weather prediction
    7. Compare ML prediction with CanSAT measurements
    8. Calculate MAE / RMSE
    9. Save ML-vs-CanSAT results
===============================================================
"""

import os
import csv
import json
import pickle
import threading
import time
from datetime import datetime

import serial
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.animation import FuncAnimation
from openpyxl import Workbook, load_workbook


# =============================================================
# CONFIGURATION
# =============================================================

SERIAL_PORT = "COM4"
BAUD_RATE = 115200

WEATHER_XLSX = "CanSAT_Weather_Merged_2018_2026.xlsx"

OUTPUT_EXCEL = "CanSAT_ML_Ground_Station.xlsx"

MODEL_PATH = "lstm_weather_model.h5"
RF_PATH = "rf_condition_model.pkl"

SCALER_X_PATH = "scaler_X.pkl"
SCALER_Y_PATH = "scaler_y.pkl"

LABEL_ENCODER_PATH = "label_encoder_conditions.pkl"

FEATURES_PATH = "numeric_features.json"


# =============================================================
# CURRENT CAN-SAT PACKET FORMAT
# =============================================================

TELEMETRY_FIELDS = [
    "packetID",
    "ms",
    "lat",
    "lon",
    "gpsAlt",
    "sats",
    "gpsFix",
    "bmpTemp",
    "bmpPress",
    "bmpAlt",
    "dhtTemp",
    "dhtHum",
    "ax",
    "ay",
    "az",
    "gx",
    "gy",
    "gz",
    "imuTemp",
    "battV",
    "rssi"
]


# =============================================================
# GLOBAL DATA
# =============================================================

data_lock = threading.Lock()

telemetry_data = []

running = True

latest_prediction = {
    "temperature": None,
    "humidity": None,
    "condition": None
}

# Keep only recent points for live graphs
MAX_GRAPH_POINTS = 300

graph_data = {
    "time": [],
    "gps_alt": [],
    "bmp_alt": [],
    "bmp_temp": [],
    "dht_temp": [],
    "dht_hum": [],
    "battery": [],
    "rssi": []
}


# =============================================================
# ML ARTIFACT LOADING
# =============================================================

def load_ml_artifacts():

    print("\n==========================================")
    print("Loading ML models...")
    print("==========================================")

    try:

        from tensorflow.keras.models import load_model

        model = load_model(
            MODEL_PATH,
            compile=False
        )

        with open(RF_PATH, "rb") as f:
            rf_model = pickle.load(f)

        with open(SCALER_X_PATH, "rb") as f:
            scaler_X = pickle.load(f)

        with open(SCALER_Y_PATH, "rb") as f:
            scaler_y = pickle.load(f)

        with open(LABEL_ENCODER_PATH, "rb") as f:
            label_encoder = pickle.load(f)

        with open(FEATURES_PATH, "r") as f:
            numeric_features = json.load(f)

        print("LSTM model              : OK")
        print("Random Forest model     : OK")
        print("Input scaler            : OK")
        print("Output scaler           : OK")
        print("Label encoder           : OK")
        print("Feature configuration   : OK")

        return (
            model,
            rf_model,
            scaler_X,
            scaler_y,
            label_encoder,
            numeric_features
        )

    except Exception as e:

        print("\nML MODEL LOADING ERROR")
        print("------------------------------------------")
        print(e)
        print("------------------------------------------")

        return None, None, None, None, None, None


# =============================================================
# PREPARE HISTORICAL DATA
# =============================================================

def prepare_historical_data(numeric_features):

    print("\nLoading historical weather dataset...")

    df = pd.read_excel(
        WEATHER_XLSX,
        sheet_name="Merged_Data"
    )

    # Existing integration expects datetime/date
    if "datetime" in df.columns:
        df.rename(
            columns={"datetime": "date"},
            inplace=True
        )

    if "date" not in df.columns:
        raise ValueError(
            "Could not find 'date' or 'datetime' column."
        )

    df["date"] = pd.to_datetime(
        df["date"],
        errors="coerce"
    )

    # Calendar features
    df["month"] = df["date"].dt.month

    df["dayofweek"] = df["date"].dt.dayofweek

    df["month_sin"] = np.sin(
        2 * np.pi * df["month"] / 12
    )

    df["month_cos"] = np.cos(
        2 * np.pi * df["month"] / 12
    )

    df["dow_sin"] = np.sin(
        2 * np.pi * df["dayofweek"] / 7
    )

    df["dow_cos"] = np.cos(
        2 * np.pi * df["dayofweek"] / 7
    )

    # Convert ML features to numeric
    for col in numeric_features:

        if col not in df.columns:

            print(
                f"WARNING: Feature '{col}' "
                f"not found in dataset."
            )

        else:

            df[col] = pd.to_numeric(
                df[col],
                errors="coerce"
            )

    # Fill missing values
    df[numeric_features] = (
        df[numeric_features]
        .ffill()
        .bfill()
    )

    print(
        f"Historical records loaded: {len(df)}"
    )

    return df


# =============================================================
# GENERATE ML FORECAST
# =============================================================

def generate_ml_forecast(
    model,
    rf_model,
    scaler_X,
    scaler_y,
    label_encoder,
    numeric_features,
    historical_data
):

    try:

        # Use latest 7 days
        latest = historical_data[
            numeric_features
        ].tail(7)

        scaled = scaler_X.transform(
            latest
        )

        sequence = scaled.reshape(
            1,
            7,
            len(numeric_features)
        )

        # -----------------------------
        # LSTM prediction
        # -----------------------------

        prediction_scaled = model.predict(
            sequence,
            verbose=0
        )

        prediction = scaler_y.inverse_transform(
            prediction_scaled
        )[0]

        predicted_temperature = float(
            prediction[0]
        )

        predicted_humidity = float(
            prediction[1]
        )

        # -----------------------------
        # Random Forest
        # -----------------------------

        condition_encoded = rf_model.predict(
            scaled[-1].reshape(1, -1)
        )

        predicted_condition = (
            label_encoder
            .inverse_transform(
                condition_encoded
            )[0]
        )

        latest_prediction[
            "temperature"
        ] = predicted_temperature

        latest_prediction[
            "humidity"
        ] = predicted_humidity

        latest_prediction[
            "condition"
        ] = str(predicted_condition)

        print("\n==========================================")
        print("             ML WEATHER FORECAST")
        print("==========================================")

        print(
            f"Predicted Temperature : "
            f"{predicted_temperature:.2f} °C"
        )

        print(
            f"Predicted Humidity    : "
            f"{predicted_humidity:.2f} %"
        )

        print(
            f"Predicted Condition   : "
            f"{predicted_condition}"
        )

        print("==========================================")

    except Exception as e:

        print("\nML prediction error:")
        print(e)


# =============================================================
# CREATE EXCEL WORKBOOK
# =============================================================

def initialize_excel():

    if os.path.exists(OUTPUT_EXCEL):

        wb = load_workbook(
            OUTPUT_EXCEL
        )

    else:

        wb = Workbook()

        # Remove default sheet
        default_sheet = wb.active
        wb.remove(default_sheet)

    # -----------------------------
    # Telemetry sheet
    # -----------------------------

    if "CanSAT_Telemetry" not in wb.sheetnames:

        ws = wb.create_sheet(
            "CanSAT_Telemetry"
        )

        ws.append(
            ["Receive Time"]
            + TELEMETRY_FIELDS
        )

    # -----------------------------
    # ML validation sheet
    # -----------------------------

    if "ML_vs_CanSAT" not in wb.sheetnames:

        ws = wb.create_sheet(
            "ML_vs_CanSAT"
        )

        ws.append([
            "Receive Time",
            "Packet ID",
            "Altitude",

            "ML Temperature",
            "CanSAT Temperature",
            "Temperature Error",

            "ML Humidity",
            "CanSAT Humidity",
            "Humidity Error",

            "ML Condition"
        ])

    # -----------------------------
    # Performance sheet
    # -----------------------------

    if "Performance" not in wb.sheetnames:

        ws = wb.create_sheet(
            "Performance"
        )

        ws.append([
            "Metric",
            "Temperature",
            "Humidity"
        ])

        ws.append([
            "MAE",
            "",
            ""
        ])

        ws.append([
            "RMSE",
            "",
            ""
        ])

        ws.append([
            "Maximum Error",
            "",
            ""
        ])

    wb.save(
        OUTPUT_EXCEL
    )

    return wb


# =============================================================
# SAVE TELEMETRY TO EXCEL
# =============================================================

def save_telemetry_to_excel(row):

    try:

        wb = load_workbook(
            OUTPUT_EXCEL
        )

        ws = wb[
            "CanSAT_Telemetry"
        ]

        receive_time = datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )

        ws.append(
            [receive_time] + row
        )

        wb.save(
            OUTPUT_EXCEL
        )

    except PermissionError:

        print(
            "\nWARNING: Excel file is open."
        )

        print(
            "Close the Excel file so Python "
            "can continue writing."
        )

    except Exception as e:

        print(
            "Excel save error:",
            e
        )


# =============================================================
# SAVE ML COMPARISON
# =============================================================

def save_ml_comparison(row):

    try:

        bmp_temp = float(row[7])

        dht_hum = float(row[11])

        altitude = float(row[9])

        packet_id = row[0]

        predicted_temp = (
            latest_prediction["temperature"]
        )

        predicted_hum = (
            latest_prediction["humidity"]
        )

        predicted_condition = (
            latest_prediction["condition"]
        )

        if predicted_temp is None:
            return

        temp_error = abs(
            predicted_temp - bmp_temp
        )

        hum_error = abs(
            predicted_hum - dht_hum
        )

        wb = load_workbook(
            OUTPUT_EXCEL
        )

        ws = wb[
            "ML_vs_CanSAT"
        ]

        ws.append([
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            ),

            packet_id,

            altitude,

            round(
                predicted_temp,
                3
            ),

            round(
                bmp_temp,
                3
            ),

            round(
                temp_error,
                3
            ),

            round(
                predicted_hum,
                3
            ),

            round(
                dht_hum,
                3
            ),

            round(
                hum_error,
                3
            ),

            predicted_condition
        ])

        wb.save(
            OUTPUT_EXCEL
        )

    except Exception as e:

        print(
            "ML comparison save error:",
            e
        )


# =============================================================
# UPDATE PERFORMANCE METRICS
# =============================================================

def update_performance():

    try:

        df = pd.read_excel(
            OUTPUT_EXCEL,
            sheet_name="ML_vs_CanSAT"
        )

        if len(df) < 1:
            return

        temp_errors = pd.to_numeric(
            df["Temperature Error"],
            errors="coerce"
        ).dropna()

        hum_errors = pd.to_numeric(
            df["Humidity Error"],
            errors="coerce"
        ).dropna()

        if len(temp_errors) == 0:
            return

        temp_mae = temp_errors.mean()

        temp_rmse = np.sqrt(
            np.mean(
                temp_errors ** 2
            )
        )

        temp_max = temp_errors.max()

        hum_mae = (
            hum_errors.mean()
            if len(hum_errors)
            else np.nan
        )

        hum_rmse = (
            np.sqrt(
                np.mean(
                    hum_errors ** 2
                )
            )
            if len(hum_errors)
            else np.nan
        )

        hum_max = (
            hum_errors.max()
            if len(hum_errors)
            else np.nan
        )

        wb = load_workbook(
            OUTPUT_EXCEL
        )

        ws = wb[
            "Performance"
        ]

        ws["B2"] = round(
            temp_mae,
            3
        )

        ws["C2"] = round(
            hum_mae,
            3
        )

        ws["B3"] = round(
            temp_rmse,
            3
        )

        ws["C3"] = round(
            hum_rmse,
            3
        )

        ws["B4"] = round(
            temp_max,
            3
        )

        ws["C4"] = round(
            hum_max,
            3
        )

        wb.save(
            OUTPUT_EXCEL
        )

    except Exception as e:

        print(
            "Performance calculation error:",
            e
        )


# =============================================================
# SERIAL RECEIVER THREAD
# =============================================================

def serial_receiver():

    global running

    print("\n==========================================")
    print("        CONNECTING TO GROUND STATION")
    print("==========================================")

    try:

        ser = serial.Serial(
            SERIAL_PORT,
            BAUD_RATE,
            timeout=1
        )

        time.sleep(2)

        print(
            f"Connected to {SERIAL_PORT}"
        )

        print(
            f"Baud rate: {BAUD_RATE}"
        )

        print(
            "Waiting for CanSAT telemetry..."
        )

        while running:

            raw = (
                ser.readline()
                .decode(
                    "utf-8",
                    errors="ignore"
                )
                .strip()
            )

            if not raw:
                continue

            # Ignore ESP32 status messages
            if (
                raw.startswith("GROUND")
                or
                raw.startswith("Waiting")
                or
                raw.startswith("packetID")
            ):
                continue

            parts = raw.split(",")

            if len(parts) != len(
                TELEMETRY_FIELDS
            ):

                print(
                    "[Skipped malformed packet]",
                    raw
                )

                continue

            try:

                # Verify packet ID
                int(parts[0])

                with data_lock:

                    telemetry_data.append(
                        parts
                    )

                    # Keep memory reasonable
                    if len(telemetry_data) > 5000:

                        telemetry_data.pop(
                            0
                        )

                update_graph_data(
                    parts
                )

                save_telemetry_to_excel(
                    parts
                )

                save_ml_comparison(
                    parts
                )

                update_performance()

                print(
                    f"Packet {parts[0]} | "
                    f"Altitude: {parts[9]} m | "
                    f"Temp: {parts[7]} °C | "
                    f"Humidity: {parts[11]} % | "
                    f"Battery: {parts[19]} V | "
                    f"RSSI: {parts[20]} dBm"
                )

            except Exception as e:

                print(
                    "Packet processing error:",
                    e
                )

        ser.close()

    except serial.SerialException as e:

        print("\nSERIAL CONNECTION ERROR")
        print("------------------------------------------")
        print(e)
        print("------------------------------------------")
        print(
            f"Check that {SERIAL_PORT} is correct "
            f"and Serial Monitor is CLOSED."
        )

    except Exception as e:

        print(
            "Receiver error:",
            e
        )


# =============================================================
# UPDATE GRAPH DATA
# =============================================================

def update_graph_data(row):

    try:

        packet_id = int(row[0])

        gps_alt = float(row[4])

        bmp_alt = float(row[9])

        bmp_temp = float(row[7])

        dht_temp = float(row[10])

        dht_hum = float(row[11])

        battery = float(row[19])

        rssi = float(row[20])

        graph_data["time"].append(
            packet_id
        )

        graph_data["gps_alt"].append(
            gps_alt
        )

        graph_data["bmp_alt"].append(
            bmp_alt
        )

        graph_data["bmp_temp"].append(
            bmp_temp
        )

        graph_data["dht_temp"].append(
            dht_temp
        )

        graph_data["dht_hum"].append(
            dht_hum
        )

        graph_data["battery"].append(
            battery
        )

        graph_data["rssi"].append(
            rssi
        )

        # Keep latest points
        for key in graph_data:

            if len(
                graph_data[key]
            ) > MAX_GRAPH_POINTS:

                graph_data[key].pop(0)

    except Exception as e:

        print(
            "Graph data error:",
            e
        )


# =============================================================
# CREATE LIVE DASHBOARD
# =============================================================

def create_dashboard():

    plt.ion()

    fig = plt.figure(
        figsize=(14, 8)
    )

    fig.suptitle(
        "CanSAT + Machine Learning Ground Station",
        fontsize=16
    )

    ax1 = fig.add_subplot(231)
    ax2 = fig.add_subplot(232)
    ax3 = fig.add_subplot(233)
    ax4 = fig.add_subplot(234)
    ax5 = fig.add_subplot(235)
    ax6 = fig.add_subplot(236)

    def update(_):

        with data_lock:

            t = list(
                graph_data["time"]
            )

            gps_alt = list(
                graph_data["gps_alt"]
            )

            bmp_alt = list(
                graph_data["bmp_alt"]
            )

            bmp_temp = list(
                graph_data["bmp_temp"]
            )

            dht_temp = list(
                graph_data["dht_temp"]
            )

            dht_hum = list(
                graph_data["dht_hum"]
            )

            battery = list(
                graph_data["battery"]
            )

            rssi = list(
                graph_data["rssi"]
            )

        # -------------------------
        # ALTITUDE
        # -------------------------

        ax1.clear()

        ax1.plot(
            t,
            gps_alt,
            label="GPS Altitude"
        )

        ax1.plot(
            t,
            bmp_alt,
            label="BMP Altitude"
        )

        ax1.set_title(
            "Altitude"
        )

        ax1.set_xlabel(
            "Packet"
        )

        ax1.set_ylabel(
            "Altitude (m)"
        )

        ax1.grid(
            True
        )

        ax1.legend()

        # -------------------------
        # TEMPERATURE
        # -------------------------

        ax2.clear()

        ax2.plot(
            t,
            bmp_temp,
            label="BMP280"
        )

        ax2.plot(
            t,
            dht_temp,
            label="DHT11"
        )

        # ML prediction line
        if (
            latest_prediction[
                "temperature"
            ] is not None
        ):

            ax2.axhline(
                latest_prediction[
                    "temperature"
                ],
                linestyle="--",
                label="ML Prediction"
            )

        ax2.set_title(
            "Temperature"
        )

        ax2.set_xlabel(
            "Packet"
        )

        ax2.set_ylabel(
            "Temperature (°C)"
        )

        ax2.grid(
            True
        )

        ax2.legend()

        # -------------------------
        # HUMIDITY
        # -------------------------

        ax3.clear()

        ax3.plot(
            t,
            dht_hum,
            label="DHT11 Humidity"
        )

        if (
            latest_prediction[
                "humidity"
            ] is not None
        ):

            ax3.axhline(
                latest_prediction[
                    "humidity"
                ],
                linestyle="--",
                label="ML Prediction"
            )

        ax3.set_title(
            "Humidity"
        )

        ax3.set_xlabel(
            "Packet"
        )

        ax3.set_ylabel(
            "Humidity (%)"
        )

        ax3.grid(
            True
        )

        ax3.legend()

        # -------------------------
        # BATTERY
        # -------------------------

        ax4.clear()

        ax4.plot(
            t,
            battery
        )

        ax4.set_title(
            "Battery Voltage"
        )

        ax4.set_xlabel(
            "Packet"
        )

        ax4.set_ylabel(
            "Voltage (V)"
        )

        ax4.grid(
            True
        )

        # -------------------------
        # RSSI
        # -------------------------

        ax5.clear()

        ax5.plot(
            t,
            rssi
        )

        ax5.set_title(
            "LoRa RSSI"
        )

        ax5.set_xlabel(
            "Packet"
        )

        ax5.set_ylabel(
            "RSSI (dBm)"
        )

        ax5.grid(
            True
        )

        # -------------------------
        # STATUS
        # -------------------------

        ax6.clear()

        ax6.axis(
            "off"
        )

        status = (
            "CanSAT ML STATUS\n\n"
        )

        if len(t) > 0:

            status += (
                f"Packets: {len(t)}\n"
            )

            status += (
                f"Altitude: "
                f"{bmp_alt[-1]:.2f} m\n"
            )

            status += (
                f"Temperature: "
                f"{bmp_temp[-1]:.2f} °C\n"
            )

            status += (
                f"Humidity: "
                f"{dht_hum[-1]:.2f} %\n"
            )

            status += (
                f"Battery: "
                f"{battery[-1]:.2f} V\n"
            )

            status += (
                f"RSSI: "
                f"{rssi[-1]:.0f} dBm\n\n"
            )

        status += (
            "ML FORECAST\n\n"
        )

        if latest_prediction[
            "temperature"
        ] is not None:

            status += (
                f"Temperature: "
                f"{latest_prediction['temperature']:.2f} °C\n"
            )

            status += (
                f"Humidity: "
                f"{latest_prediction['humidity']:.2f} %\n"
            )

            status += (
                f"Condition: "
                f"{latest_prediction['condition']}\n"
            )

        else:

            status += (
                "ML prediction unavailable"
            )

        ax6.text(
            0.05,
            0.95,
            status,
            transform=ax6.transAxes,
            verticalalignment="top",
            fontsize=11
        )

        fig.tight_layout()

    animation = FuncAnimation(
        fig,
        update,
        interval=1000,
        cache_frame_data=False
    )

    plt.show(
        block=True
    )


# =============================================================
# MAIN
# =============================================================

def main():

    global running

    print("\n")
    print("====================================================")
    print("        CanSAT ML GROUND STATION")
    print("====================================================")

    print(
        f"Serial Port : {SERIAL_PORT}"
    )

    print(
        f"Baud Rate   : {BAUD_RATE}"
    )

    print(
        f"Excel File  : {OUTPUT_EXCEL}"
    )

    print("====================================================")

    # -----------------------------------------
    # Initialize Excel
    # -----------------------------------------

    initialize_excel()

    # -----------------------------------------
    # Load ML
    # -----------------------------------------

    (
        model,
        rf_model,
        scaler_X,
        scaler_y,
        label_encoder,
        numeric_features
    ) = load_ml_artifacts()

    # -----------------------------------------
    # Generate initial ML forecast
    # -----------------------------------------

    if model is not None:

        try:

            historical_data = (
                prepare_historical_data(
                    numeric_features
                )
            )

            generate_ml_forecast(
                model,
                rf_model,
                scaler_X,
                scaler_y,
                label_encoder,
                numeric_features,
                historical_data
            )

        except Exception as e:

            print(
                "\nCould not generate ML forecast:"
            )

            print(e)

    else:

        print(
            "\nRunning without ML prediction."
        )

    # -----------------------------------------
    # Start serial receiver
    # -----------------------------------------

    receiver_thread = threading.Thread(
        target=serial_receiver,
        daemon=True
    )

    receiver_thread.start()

    # -----------------------------------------
    # Start dashboard
    # -----------------------------------------

    try:

        create_dashboard()

    except KeyboardInterrupt:

        print(
            "\nStopping ground station..."
        )

    finally:

        running = False

        print(
            "\nGround station stopped."
        )


# =============================================================
# PROGRAM START
# =============================================================

if __name__ == "__main__":

    main()
