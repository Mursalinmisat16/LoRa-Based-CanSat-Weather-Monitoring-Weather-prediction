import serial
import openpyxl
import os
from datetime import datetime

# =========================
# SETTINGS
# =========================

SERIAL_PORT = "COM4"       # CHANGE THIS
BAUD_RATE = 115200

EXCEL_FILE = "CanSAT_Telemetry.xlsx"
SHEET_NAME = "Telemetry"

# =========================
# CSV HEADERS
# =========================

HEADERS = [
    "Receive Time",
    "packetID",
    "ms",
    "lat",
    "lon",
    "gpsAlt",
    "sats",
    "gpsFix",
    "bmpTemp",
    "bmpPress",
    "bmpAlt",
    "dhtTemp",
    "dhtHum",
    "ax",
    "ay",
    "az",
    "gx",
    "gy",
    "gz",
    "imuTemp",
    "battV",
    "rssi"
]

# =========================
# CREATE / OPEN EXCEL
# =========================

if os.path.exists(EXCEL_FILE):
    workbook = openpyxl.load_workbook(EXCEL_FILE)
    print(f"Existing Excel file opened: {EXCEL_FILE}")

    if SHEET_NAME in workbook.sheetnames:
        sheet = workbook[SHEET_NAME]
    else:
        sheet = workbook.create_sheet(SHEET_NAME)
        sheet.append(HEADERS)

else:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(HEADERS)

    workbook.save(EXCEL_FILE)

    print(f"New Excel file created: {EXCEL_FILE}")


# =========================
# CONNECT TO ESP32
# =========================

print()
print("Connecting to ESP32...")
print(f"Port: {SERIAL_PORT}")
print(f"Baud rate: {BAUD_RATE}")

try:
    ser = serial.Serial(
        SERIAL_PORT,
        BAUD_RATE,
        timeout=1
    )

except Exception as e:
    print()
    print("ERROR: Could not open serial port.")
    print(e)
    print()
    print("Check:")
    print("1. ESP32 is connected")
    print("2. Correct COM port is selected")
    print("3. Arduino Serial Monitor is CLOSED")
    exit()


print()
print("================================")
print("   CANSAT EXCEL LOGGER STARTED")
print("================================")
print()
print("Waiting for telemetry...")
print("Press CTRL+C to stop.")
print()


# =========================
# READ SERIAL DATA
# =========================

try:

    while True:

        line = ser.readline().decode(
            "utf-8",
            errors="ignore"
        ).strip()

        if not line:
            continue

        # Print everything received
        print(line)

        # Ignore status messages
        if line.startswith("GROUND"):
            continue

        if line.startswith("Waiting"):
            continue

        # Ignore CSV header
        if line.startswith("packetID"):
            continue

        # Only process telemetry packets
        values = line.split(",")

        # Expected:
        # packetID through RSSI = 21 values

        if len(values) != 21:
            print(
                f"Skipping invalid packet: "
                f"{len(values)} fields"
            )
            continue

        # Add PC receive time
        receive_time = datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )

        row = [receive_time] + values

        # Write to Excel
        sheet.append(row)

        # Save immediately
        workbook.save(EXCEL_FILE)

        print(
            f"Saved packet {values[0]} "
            f"| RSSI: {values[-1]} dBm"
        )


except KeyboardInterrupt:

    print()
    print("Stopping logger...")

finally:

    ser.close()
    workbook.save(EXCEL_FILE)

    print("Excel file saved.")
    print("Serial connection closed.")